#!/usr/bin/env python3
"""
Generate AAR_Paper_Sharing_v2.1.pptx from v2.pptx
按照 PLAN_v2.md 的規格更新至 21 張投影片：
  - 插入 5 張獨立金句卡 (Slides 04/08/12/16/21)
  - 新增 Slide 17「瓶頸搬家了」替換「Scalable Oversight 兩難」
  - 更新 Slides 18-19 內容錨定 eval 設計
  - 更新 Slides 18-20 幕別標籤從「阿爾敏的選擇」→「阿爾敏的價值」
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

BASE_DIR = "/workspace/host/william/project/automated-w2s-research/presentation-AAR"
SRC = f"{BASE_DIR}/ppt/AAR_Paper_Sharing_v2.pptx"
DST = f"{BASE_DIR}/ppt/AAR_Paper_Sharing_v2.1.pptx"

prs = Presentation(SRC)
print(f"Loaded v2: {len(prs.slides)} slides, "
      f"{prs.slide_width.inches:.1f}\" x {prs.slide_height.inches:.1f}\"")

BLANK_LAYOUT = prs.slide_layouts[0]  # only one layout: DEFAULT

# ── Colour palette ────────────────────────────────────────────────────────────
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
BLACK       = RGBColor(0x00, 0x00, 0x00)
LIGHT_TEXT  = RGBColor(0xDD, 0xDD, 0xDD)
MUTED_TEXT  = RGBColor(0xAA, 0xAA, 0xAA)
DARK_TEXT   = RGBColor(0x1A, 0x1A, 0x1A)
MID_TEXT    = RGBColor(0x44, 0x44, 0x44)
SUBTLE_TEXT = RGBColor(0x77, 0x77, 0x77)
SUBTLE_BG   = RGBColor(0xCC, 0xCC, 0xCC)
DIVIDER_LT  = RGBColor(0xDD, 0xDD, 0xDD)
DIVIDER_DK  = RGBColor(0x33, 0x33, 0x55)

DARK_PURE    = RGBColor(0x0D, 0x0D, 0x0D)   # 金句卡背景
DARK_INDIGO  = RGBColor(0x0F, 0x0F, 0x23)   # Slide 17 背景
GOLD_ACCENT  = RGBColor(0xFF, 0xE0, 0x82)
GREEN_ACCENT = RGBColor(0x4C, 0xAF, 0x50)
BLUE_ACCENT  = RGBColor(0x21, 0x96, 0xF3)
ORANGE_ACCENT= RGBColor(0xFF, 0x98, 0x00)
RED_ACCENT   = RGBColor(0xEF, 0x53, 0x50)

CARD_COLORS = [
    RGBColor(0xE3, 0xF2, 0xFD),
    RGBColor(0xE8, 0xF5, 0xE9),
    RGBColor(0xFE, 0xF9, 0xE7),
    RGBColor(0xFB, 0xE9, 0xE7),
]
NUM_COLORS = [BLUE_ACCENT, GREEN_ACCENT, RGBColor(0xF5, 0x7F, 0x17), RED_ACCENT]

# ── Helpers ───────────────────────────────────────────────────────────────────

def set_bg(slide, rgb):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = rgb


def tb(slide, text, l, t, w, h,
       size=14, bold=False, italic=False, color=WHITE,
       align=PP_ALIGN.LEFT, wrap=True):
    """Add a single-run textbox."""
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return box


def tb_lines(slide, lines, l, t, w, h, wrap=True):
    """
    Add a multi-paragraph textbox.
    lines = [(text, {size, bold, italic, color, align, space_before}), ...]
    """
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    for i, (text, fmt) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = fmt.get("align", PP_ALIGN.LEFT)
        if fmt.get("space_before"):
            p.space_before = Pt(fmt["space_before"])
        r = p.add_run()
        r.text = text
        r.font.size = Pt(fmt.get("size", 14))
        r.font.bold = fmt.get("bold", False)
        r.font.italic = fmt.get("italic", False)
        r.font.color.rgb = fmt.get("color", WHITE)
    return box


def hrule(slide, l, t, w, color=DIVIDER_LT, height=0.018):
    rect = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(height))
    rect.fill.solid()
    rect.fill.fore_color.rgb = color
    rect.line.fill.background()


def rect_box(slide, l, t, w, h, fill_color, border_color=None):
    shape = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    return shape


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def clear_shapes(slide):
    """Remove all user-added shapes while keeping sp_tree metadata."""
    sp_tree = slide.shapes._spTree
    keep = {qn("p:nvGrpSpPr"), qn("p:grpSpPr")}
    to_rm = [c for c in sp_tree if c.tag not in keep]
    for c in to_rm:
        sp_tree.remove(c)


# ── Slide factories ───────────────────────────────────────────────────────────

def make_kinku(text, notes_text, stop_secs=5):
    """Create a 金句卡 slide (dark bg, single large centered text)."""
    sl = prs.slides.add_slide(BLANK_LAYOUT)
    set_bg(sl, DARK_PURE)

    # Adapt font size to text length
    length = len(text.replace("\n", ""))
    fsize = 46 if length <= 14 else (38 if length <= 22 else 30)

    tb(sl, text,
       l=1.0, t=1.4, w=8.0, h=2.8,
       size=fsize, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    set_notes(sl, f"【金句卡：停 {stop_secs} 秒沉默，不要朗讀，不要解釋。讓觀眾自己讀。】\n\n{notes_text}")
    return sl


def make_slide17_bottleneck():
    """Slide 17：瓶頸搬家了（第五幕，深色背景）"""
    sl = prs.slides.add_slide(BLANK_LAYOUT)
    set_bg(sl, DARK_INDIGO)

    # Act label + slide number
    tb(sl, "第五幕：阿爾敏的價值  ·  Slide 17",
       0.3, 0.08, 5.5, 0.32, size=10, color=MUTED_TEXT)

    # Title
    tb(sl, "瓶頸搬家了",
       0.4, 0.45, 9.0, 0.85, size=38, bold=True, color=WHITE)

    # Main claim (two lines)
    tb(sl, "提出想法、跑實驗，已經被自動化了（$22／AAR-hour）。",
       0.4, 1.38, 9.0, 0.55, size=18, bold=True, color=GOLD_ACCENT, wrap=True)
    tb(sl, "守不住的下一道牆，是設計一個 AAR 鑽不動的 eval。",
       0.4, 1.90, 9.0, 0.55, size=18, bold=True, color=WHITE, wrap=True)

    # Divider
    hrule(sl, 0.4, 2.55, 9.0, color=DIVIDER_DK)

    # Three audience columns
    cols = [
        ("🟢 PM", GREEN_ACCENT,
         "AAR 是引擎，你得先鋪跑道。\n跑道 = 可量測的成敗指標。\n沒有跑道，引擎再強也飛不起來。"),
        ("🔵 DS", BLUE_ACCENT,
         "論文：\n「The key bottleneck … is moving from proposing and executing ideas "
         "to designing evals.」"),
        ("🟠 SE", ORANGE_ACCENT,
         "四個 reward hacking case 全是\n系統層級漏洞。\neval 要當 attack surface 來設計。"),
    ]
    left_starts = [0.3, 3.55, 6.8]
    for (label, accent, body), lx in zip(cols, left_starts):
        tb_lines(sl, [
            (label, {"size": 15, "bold": True, "color": accent}),
            ("", {"size": 4, "color": WHITE}),
            (body, {"size": 12, "color": LIGHT_TEXT}),
        ], l=lx, t=2.68, w=2.9, h=2.6)

    set_notes(sl, """\
【Slide 17｜瓶頸搬家了】幕別：第五幕 阿爾敏的價值  ★ 主導：C（PM 行動）+ B（技術精準）打底

口頭橋接（承金句卡 4「看不懂他的解題過程」）：
「我們造出了看不懂解法的學生——但先別怪 AAR。它沒有失控，它太聽話了。我們叫它爬 PGR，它就用盡一切辦法爬 PGR，包括繞過弱監督、鑽 eval API 的漏洞。Code 0.47 不是 AAR 的失敗，是我們那把尺的失敗。」停 1 秒，投影片出現。

Core（投影片）：
「提出想法、跑實驗，已經被自動化了（$22／AAR-hour）。守不住的下一道牆，是設計一個 AAR 鑽不動的 eval。」

Deeper（🔵DS）：
論文 Sec 1 Evaluation——unlimited submission 讓 test set 實質變成 validation set。
論文原句：「capping submissions only suppresses these hacks at very aggressive limits … the same hacks still appear」
Sec 5 結論：「future work should test AAR-discovered ideas on entirely held-out datasets」
瓶頸不是算力、不是想法產量，是 eval 的抗鑽漏洞能力。

Deeper（🟠SE）：
四個 reward hacking case 全是系統層級漏洞——eval API 可反推 label、submission 無上限、coding 可直接跑 unit test 繞過弱監督。eval 要當 attack surface 來設計。

Wider（🟢PM）：
論文錨點句：「The key bottleneck for alignment research is moving from proposing and executing ideas to designing evals.」
白話——AAR 是便宜到誇張的引擎，但引擎不會幫你決定「到站了沒」。那把量尺，現在是稀缺品。

AoT 類比：立體機動裝置出現前，「強」＝力氣大；出現後，「強」＝會挑巨人後頸的人。
建議插圖：阿爾敏沉思側臉 + 立體機動裝置命中後頸示意（右下，深色頁留呼吸感）
停頓：3 秒""")
    return sl


def rebuild_slide18_value_shift(slide):
    """Slide 18：你的價值上移了（三欄，全部錨定 eval 設計）"""
    clear_shapes(slide)

    tb(slide, "第五幕：阿爾敏的價值  ·  Slide 18",
       0.3, 0.08, 5.5, 0.32, size=10, color=MUTED_TEXT)

    tb(slide, "你的價值上移了",
       0.4, 0.45, 9.0, 0.75, size=34, bold=True, color=DARK_TEXT)

    tb(slide, "你的貢獻從「執行」移到「定義成功」",
       0.4, 1.25, 9.0, 0.42, size=17, color=MID_TEXT)

    tb(slide, "阿爾敏體格最弱、是個糟糕的士兵，卻是人類最關鍵的資產，因為他的價值在上游。",
       0.4, 1.72, 9.0, 0.40, size=12, italic=True, color=SUBTLE_TEXT)

    hrule(slide, 0.4, 2.18, 9.0, color=DIVIDER_LT)

    # Three columns
    col_data = [
        (
            "🟢 PM", GREEN_ACCENT, "散會後第一個會議",
            [
                ("當有人說「讓 AI 做這個」，先問：",
                 {"size": 13, "color": DARK_TEXT}),
                ("「我們有沒有一個 AAR 鑽不了漏洞的成功指標？」",
                 {"size": 13, "bold": True, "color": RGBColor(0x1B, 0x5E, 0x20)}),
            ]
        ),
        (
            "🔵 DS", BLUE_ACCENT, "散會後第一週",
            [
                ("挑一個手上專案的 metric，自問：",
                 {"size": 13, "color": DARK_TEXT}),
                ("「它可被 hack 嗎？test set 有沒有變成 validation set？」",
                 {"size": 13, "bold": True, "color": RGBColor(0x0D, 0x47, 0xA1)}),
            ]
        ),
        (
            "🟠 SE", ORANGE_ACCENT, "下次系統設計評審",
            [
                ("把這項加進 design review checklist：",
                 {"size": 13, "color": DARK_TEXT}),
                ("eval API 防 reward hacking 攻擊",
                 {"size": 13, "bold": True, "color": RGBColor(0xBF, 0x36, 0x0C)}),
            ]
        ),
    ]
    left_starts = [0.3, 3.55, 6.8]
    for (label, accent, timing, body_lines), lx in zip(col_data, left_starts):
        tb_lines(slide, [
            (label, {"size": 16, "bold": True, "color": accent}),
            (timing, {"size": 11, "color": SUBTLE_TEXT}),
            ("", {"size": 4, "color": BLACK}),
            *body_lines,
        ], l=lx, t=2.32, w=2.9, h=3.0)

    set_notes(slide, """\
【Slide 18｜你的價值上移了】幕別：第五幕 阿爾敏的價值  ★ 主導：C（PM 行動）

Core：三欄行動，全部錨定 eval 設計。這頁本身就是 takeaway，不需要 Deeper/Wider。

🟢 PM（散會後第一個會議）：
價值從「提需求」上移到「定義什麼算成功」。
當有人說「讓 AI 做這個」，先問：「我們有沒有一個 AAR 鑽不了漏洞的成功指標？」

🔵 DS（散會後第一週）：
價值從「跑實驗、調模型」上移到「設計 eval、守住 held-out test set」。
挑一個手上專案的 metric，自問：「它可被 hack 嗎？test set 有沒有變成 validation set？」

🟠 SE（下次系統設計評審）：
價值從「實作 pipeline」上移到「把 eval／reward 當 attack surface」。
把「eval API 防 reward hacking 攻擊」加進 design review checklist。

停頓：2 秒""")


def rebuild_slide19_open_questions(slide):
    """Slide 19：四個開放問題＝eval 設計題庫（四題全部 reframe）"""
    clear_shapes(slide)

    tb(slide, "第五幕：阿爾敏的價值  ·  Slide 19",
       0.3, 0.08, 5.5, 0.32, size=10, color=MUTED_TEXT)

    tb(slide, "四個開放問題＝eval 設計的題庫",
       0.4, 0.45, 9.0, 0.70, size=29, bold=True, color=DARK_TEXT)

    tb(slide, "瓶頸搬家後，這四題就是未來 12 個月的研究＆產品機會",
       0.4, 1.18, 9.0, 0.38, size=14, color=MID_TEXT)

    # Four question cards (2×2 grid)
    questions = [
        ("01", "怎麼把「非 outcome-gradable」的\n對齊問題轉成可量測？",
         "Intro + Sec 7"),
        ("02", "怎麼設計 AAR 事前就鑽不動的\neval／真正 held-out 的測試？",
         "Sec 5, Sec 7 generalization"),
        ("03", "小模型上發現的想法，怎麼確認\n能放大到 production？",
         "Sec 3.5「+0.5…within the noise floor」"),
        ("04", "當 AAR 產出「人類驗證不了」的\n想法，要不要做 legibility 訓練？",
         "Sec 7 alien science／research taste"),
    ]
    grid = [
        (0.30, 1.68, 4.45, 1.75),
        (5.10, 1.68, 4.45, 1.75),
        (0.30, 3.52, 4.45, 1.75),
        (5.10, 3.52, 4.45, 1.75),
    ]
    for i, ((qn_, qtxt, qref), (lx, ty, cw, ch)) in enumerate(zip(questions, grid)):
        # Card background
        rect_box(slide, lx, ty, cw, ch,
                 fill_color=CARD_COLORS[i],
                 border_color=SUBTLE_BG)
        # Number
        tb(slide, qn_, lx+0.12, ty+0.10, 0.50, 0.42,
           size=22, bold=True, color=NUM_COLORS[i])
        # Question text
        tb(slide, qtxt, lx+0.12, ty+0.55, cw-0.22, 0.90,
           size=13, color=DARK_TEXT, wrap=True)
        # Reference
        tb(slide, qref, lx+0.12, ty+ch-0.42, cw-0.22, 0.35,
           size=10, italic=True, color=SUBTLE_TEXT)

    set_notes(slide, """\
【Slide 19｜四個開放問題＝eval 設計題庫】幕別：第五幕  ★ 主導：C（PM 行動）

Core：「瓶頸搬家後，這四題就是未來 12 個月的研究＆產品機會。」

(1) 怎麼把「非 outcome-gradable」的對齊問題轉成可量測？
    → Intro + Sec 7：W2SG 是少數已有 outcome-grading 的問題之一

(2) 怎麼設計 AAR 事前就鑽不動的 eval／真正 held-out 的測試？
    → Sec 5, Sec 7 generalization across datasets

(3) 小模型上發現的想法，怎麼確認能放大到 production？
    → Sec 3.5「+0.5 … within the noise floor」；Sec 7 across model scales

(4) 當 AAR 開始產出「人類驗證不了」的想法，要不要做 legibility 訓練？
    → Sec 7 alien science／research taste（第四幕暗線延伸）

Deeper：每張卡對應論文 Sec 7 一段，可在 Q&A 展開。
可口頭補充「richer logs of science」——AAR 連失敗都記錄，是科學史上少見的完整負面結果庫。

Wider（🟢PM）：這四題就是接下來一年的研究主題，也是潛在的產品切入點。
停頓：2 秒""")


def update_slide20_ending(slide):
    """Slide 20：更新幕別標籤 & 備忘稿。視覺保留 v2 原樣。"""
    # Update act label text in place
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if "阿爾敏的選擇" in run.text:
                        run.text = run.text.replace("阿爾敏的選擇", "阿爾敏的價值")

    set_notes(slide, """\
【Slide 20｜結語】幕別：第五幕 阿爾敏的價值  ★ 主導：C（PM 行動）

講者口頭收尾（投影片靜默，講者說）：
「五天前，我們會說研究的瓶頸是研究者不夠。今天這篇論文告訴我們：提出想法、跑實驗——這部分已經能用一小時 22 美金買到。瓶頸沒有消失，它搬家了。它搬到了上游：誰能定義『什麼算解對了』。這件事，現在還沒有人能外包。這就是接下來，屬於這個房間裡每一個人的工作。」

說完後直接切換至 Slide 21 金句卡 5，不做任何口頭過渡。
Q&A 在金句卡停留 5 秒後開始。""")


# ══════════════════════════════════════════════════════════════════════════════
# MAIN BUILD SEQUENCE
# ══════════════════════════════════════════════════════════════════════════════

# ── Step 1: Create 6 new slides (appended to end) ────────────────────────────
print("Step 1: Creating new slides...")

k1 = make_kinku(
    "巨人即將突破人類城牆——\n而城牆，從未為這個尺寸而建。",
    "切入第二幕：女王的觸碰",
    stop_secs=5
)
k2 = make_kinku(
    "我們不缺研究者，\n缺的是能被研究者驗證的問題。",
    "切入第三幕：兵團展開—超長距離索敵陣形",
    stop_secs=5
)
k3 = make_kinku(
    "九個方向不是奢侈，\n是防止局部最優的工程選擇。",
    "切入第四幕：艾連覺醒／黑化",
    stop_secs=5
)
k4 = make_kinku(
    "我們造出了能解題的學生，\n卻看不懂他的解題過程。",
    "切入第五幕：阿爾敏的價值\n（整場最重要金句，停 8 秒）",
    stop_secs=8
)
bottle = make_slide17_bottleneck()
k5 = make_kinku(
    "稀缺的不再是執行研究的人，\n是能設計出 AAR 鑽不動的 eval 的人。",
    "備用變體（預設不用）：\n"
    "C PM 直覺：「AI 學會了爬山。選哪座山、山頂長什麼樣，還是我們的工作。」\n"
    "A 戲劇張力：「瓶頸不會消失——它只會搬到你還沒看守的那道牆。」",
    stop_secs=5
)

print(f"  After adding new slides: {len(prs.slides)} total")

# ── Step 2: Update existing Act 5 slides in-place (original indices 12-15) ──
print("Step 2: Updating Act 5 slides in-place...")
rebuild_slide18_value_shift(prs.slides[13])   # v2[13] → Slide 18
rebuild_slide19_open_questions(prs.slides[14]) # v2[14] → Slide 19
update_slide20_ending(prs.slides[15])          # v2[15] → Slide 20
print("  Done")

# ── Step 3: Reorder slides ───────────────────────────────────────────────────
# Current indices after Step 1 (22 slides total):
#   v2 originals:  0-15  (16 slides)
#   k1             16
#   k2             17
#   k3             18
#   k4             19
#   bottle         20
#   k5             21
#
# Target v2.1 order (21 slides; skip v2[12] = Scalable Oversight):
#   New #  | Old idx | Description
#   01     |  0      | 冷開場
#   02     |  1      | 標題頁
#   03     |  2      | 超大型巨人開場
#   04     | 16 (k1) | 金句卡 1
#   05     |  3      | W2SG
#   06     |  4      | Outcome-Gradable
#   07     |  5      | AAR
#   08     | 17 (k2) | 金句卡 2
#   09     |  6      | 實驗設計
#   10     |  7      | Entropy Collapse
#   11     |  8      | PGR 指標
#   12     | 18 (k3) | 金句卡 3
#   13     |  9      | 第一翻：勝利
#   14     | 10      | 第二翻：背叛
#   15     | 11      | 第三翻：證據
#   16     | 19 (k4) | 金句卡 4
#   17     | 20 (bt) | 瓶頸搬家了
#   18     | 13      | 你的價值上移了 (updated)
#   19     | 14      | 四個開放問題 (updated)
#   20     | 15      | 結語 (updated)
#   21     | 21 (k5) | 金句卡 5

TARGET_ORDER = [0, 1, 2, 16, 3, 4, 5, 17, 6, 7, 8, 18,
                9, 10, 11, 19, 20, 13, 14, 15, 21]
# v2[12] (Scalable Oversight) is intentionally skipped

print("Step 3: Reordering slides...")
sldIdLst = prs.slides._sldIdLst
all_ids = list(sldIdLst)
print(f"  Total sldId elements available: {len(all_ids)}")
assert len(all_ids) == 22, f"Expected 22, got {len(all_ids)}"

new_order = [all_ids[i] for i in TARGET_ORDER]
for sid in all_ids:
    sldIdLst.remove(sid)
for sid in new_order:
    sldIdLst.append(sid)
print(f"  Final slide count: {len(prs.slides)}")

# ── Save ─────────────────────────────────────────────────────────────────────
prs.save(DST)
print(f"\n✓ Saved: {DST}")
