#!/usr/bin/env python3
"""
sync_notes.py — write speaker notes from slides.yaml into an existing .pptx.

Because the pptx slide order diverges from slides.yaml IDs (manually added slides,
reordering), you must supply an explicit mapping.  Visual slide content is untouched;
only the notes pane is written.

Mapping sources (highest → lowest priority):
  1. --map "yaml_id:pptx_idx,..."   CLI override
  2. pptx_index: <N>                field in slides.yaml entry (0-based)
  3. Entries without any mapping are SKIPPED (never assume id-1 = index)

Notes format written to pptx:
  [Core]
  ...

  [Deeper — DS]
  ...

  [Wider — PM]
  ...

Usage:
  # Sync all mapped slides (dry-run first):
  poetry run python code/sync_notes.py --dry-run

  # Sync specific yaml IDs with explicit index mapping:
  poetry run python code/sync_notes.py --map "14:22,15:23" --slides 14,15

  # Sync all slides that have pptx_index in slides.yaml:
  poetry run python code/sync_notes.py

  # Write to a new file instead of overwriting:
  poetry run python code/sync_notes.py --map "14:22,15:23" --out ppt/AAR_Paper_Sharing_v4.3.pptx
"""

import argparse
import os
import sys
import yaml
from pptx import Presentation

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_PPTX = os.path.join(BASE_DIR, "ppt", "AAR_Paper_Sharing_v4.2.pptx")
SLIDES_YAML  = os.path.join(BASE_DIR, "slides.yaml")


def load_yaml(path: str) -> list[dict]:
    with open(path, encoding="utf-8") as f:
        data = yaml.safe_load(f)
    return data.get("slides", [])


def build_notes_text(entry: dict) -> str:
    notes = entry.get("notes", {})
    parts = []

    core = (notes.get("core") or "").strip()
    if core:
        parts.append(f"[Core]\n{core}")

    # accept both 'deeper_ds' and legacy 'deeper'
    deeper = (notes.get("deeper_ds") or notes.get("deeper") or "").strip()
    if deeper:
        parts.append(f"[Deeper — DS]\n{deeper}")

    wider = (notes.get("wider_pm") or notes.get("wider") or "").strip()
    if wider:
        parts.append(f"[Wider — PM]\n{wider}")

    return "\n\n".join(parts)


def parse_map_arg(s: str) -> dict[int, int]:
    """Parse '14:22,15:23' → {14: 22, 15: 23}"""
    result = {}
    for pair in s.split(","):
        pair = pair.strip()
        if not pair:
            continue
        yaml_id, pptx_idx = pair.split(":")
        result[int(yaml_id.strip())] = int(pptx_idx.strip())
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--pptx",    default=DEFAULT_PPTX,
                        help="Source/target .pptx (default: ppt/AAR_Paper_Sharing_v4.2.pptx)")
    parser.add_argument("--out",     default=None,
                        help="Output path; if omitted, overwrites --pptx in-place")
    parser.add_argument("--yaml",    default=SLIDES_YAML,
                        help="slides.yaml path")
    parser.add_argument("--map",     default=None,
                        help='Explicit yaml_id:pptx_idx pairs, e.g. "14:22,15:23"')
    parser.add_argument("--slides",  default=None,
                        help="Comma-separated yaml IDs to sync, e.g. \"14,15\"; "
                             "if omitted, syncs all mapped slides")
    parser.add_argument("--dry-run", action="store_true",
                        help="Print what would change, write nothing")
    args = parser.parse_args()

    # ── Load data ───────────────────────────────────────────────────────
    slides_data = load_yaml(args.yaml)
    prs = Presentation(args.pptx)
    total_slides = len(prs.slides)

    # ── Build id → pptx_index mapping ──────────────────────────────────
    # Start from pptx_index fields in yaml, then apply CLI overrides
    id_to_idx: dict[int, int] = {}
    for entry in slides_data:
        yaml_id = entry.get("id")
        if yaml_id is None:
            continue
        pptx_idx = entry.get("pptx_index")
        if pptx_idx is not None:
            id_to_idx[yaml_id] = int(pptx_idx)

    if args.map:
        id_to_idx.update(parse_map_arg(args.map))

    # ── Filter by --slides ──────────────────────────────────────────────
    if args.slides:
        requested_ids = {int(x.strip()) for x in args.slides.split(",")}
    else:
        requested_ids = None  # all mapped

    # ── Build a lookup: yaml_id → entry ────────────────────────────────
    entry_by_id = {e["id"]: e for e in slides_data if "id" in e}

    # ── Determine work items ────────────────────────────────────────────
    work = []  # list of (yaml_id, pptx_idx, notes_text)
    for yaml_id, pptx_idx in sorted(id_to_idx.items()):
        if requested_ids is not None and yaml_id not in requested_ids:
            continue
        if yaml_id not in entry_by_id:
            print(f"  WARN  yaml id:{yaml_id} not found in slides.yaml — skipped")
            continue
        if pptx_idx < 0 or pptx_idx >= total_slides:
            print(f"  WARN  pptx index {pptx_idx} out of range (0–{total_slides-1}) "
                  f"for yaml id:{yaml_id} — skipped")
            continue
        notes_text = build_notes_text(entry_by_id[yaml_id])
        work.append((yaml_id, pptx_idx, notes_text))

    if not work:
        print("Nothing to sync.  Provide --map or add pptx_index fields to slides.yaml.")
        sys.exit(0)

    # ── Preview / apply ─────────────────────────────────────────────────
    label = "[DRY-RUN] " if args.dry_run else ""
    for yaml_id, pptx_idx, notes_text in work:
        slide = prs.slides[pptx_idx]
        try:
            current = slide.notes_slide.notes_text_frame.text
        except Exception:
            current = ""
        changed = (current.strip() != notes_text.strip())
        status = "UPDATE" if changed else "same  "
        print(f"  {label}{status}  yaml id:{yaml_id:2d} → pptx[{pptx_idx:02d}]"
              f"  ({len(notes_text)} chars)")
        if changed and args.dry_run:
            preview = notes_text[:120].replace("\n", "↵")
            print(f"           ↳ {preview}…")
        if changed and not args.dry_run:
            slide.notes_slide.notes_text_frame.text = notes_text

    # ── Save ────────────────────────────────────────────────────────────
    if args.dry_run:
        print("\n(dry-run — no file written)")
        return

    out_path = args.out or args.pptx
    prs.save(out_path)
    print(f"\nSaved → {out_path}  ({len(work)} slide(s) synced)")


if __name__ == "__main__":
    main()
